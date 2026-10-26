#!/usr/bin/env python3
"""
Nạp dữ liệu giả lập Power BI (MamVan_DuLieu_GiaLap_10HocSinh_v3-2.xlsx) vào DB Mầm Văn.

Chạy SAU khi: DB đã xóa sạch (GĐ2), migration_008 đã chạy (GĐ3), tài khoản
giáo viên chủ gv001 đã tồn tại. Script từ chối chạy nếu các bảng đích chưa trống.

Quyết định của chủ dự án (21/10/2026):
- Dữ liệu thuộc giáo viên gv001; mật khẩu khởi tạo của 10 học sinh = mật khẩu mặc định hệ thống.
- SHORT_ANSWER gộp vào essay (hệ thống giữ đúng 4 dạng câu hỏi).
- Đáp án đúng suy ra từ chính dữ liệu file (câu AUTO có đúng 1 đáp án được chấm đúng) -> tích sẵn.
- dim_date nạp thành bảng thật.

Dùng: python3 import_bi_mock_v3.py [--verify]
Kết nối: đọc mật khẩu DB từ tài liệu hướng dẫn nội bộ, qua tunnel 127.0.0.1:15432.
"""
import json
import re
import ssl
import sys
import uuid
from datetime import datetime, date, time, timedelta, timezone

import pg8000
from openpyxl import load_workbook

XLSX = '/home/hatch/workspace/user/files/MamVan_DuLieu_GiaLap_10HocSinh_v3-2.xlsx'
GUIDE = '/home/hatch/workspace/user/files/HUONG_DAN_CAI_DAT_VA_QUY_TRINH_MAY_MOI.md'
TZ7 = timezone(timedelta(hours=7))
DEFAULT_STUDENT_PASSWORD = '123456'

TYPE_MAP = {
    'SINGLE_CHOICE': 'single_choice',
    'MULTIPLE_CHOICE': 'multiple_choice',
    'FILL_IN': 'fill_in_blank',
    'SHORT_ANSWER': 'essay',
    'ESSAY': 'essay',
}
SHORT_TYPE = {'single_choice': 'single', 'multiple_choice': 'multi', 'fill_in_blank': 'fill', 'essay': 'essay'}
DIFF_APP = {'EASY': 'DE', 'MEDIUM': 'TB', 'HARD': 'KHO'}


def connect():
    guide = open(GUIDE, encoding='utf-8').read()
    pw = None
    for line in guide.splitlines():
        if 'Database Password' in line:
            m = re.search(r'Database Password[:\s*]*`?([^\s`*]+)`?', line)
            if m:
                pw = m.group(1)
                break
    assert pw, 'khong tim thay mat khau DB'
    ctx = ssl.create_default_context()
    ctx.check_hostname = False
    ctx.verify_mode = ssl.CERT_NONE
    return pg8000.connect(host='127.0.0.1', port=15432, user='postgres',
                          password=pw, database='postgres', ssl_context=ctx)


def sheet_rows(wb, name):
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h).strip() for h in rows[0]]
    out = []
    for r in rows[1:]:
        if all(v is None for v in r):
            continue
        out.append({hdr[i]: r[i] for i in range(len(hdr))})
    return out


def ts(v):
    """Chuẩn hoá mốc thời gian trong file (giờ VN) -> ISO có +07."""
    if v is None:
        return None
    if isinstance(v, datetime):
        if v.tzinfo is None:
            v = v.replace(tzinfo=TZ7)
        return v.isoformat()
    if isinstance(v, date):
        return datetime.combine(v, time(0, 0), TZ7).isoformat()
    return str(v)


def dt_only(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return str(v)[:10]


def boolean(v):
    if isinstance(v, bool):
        return v
    if v is None:
        return None
    return str(v).strip().upper() in ('TRUE', '1', 'YES')


def num(v):
    if v is None or v == '':
        return None
    return float(v)


def derive_answers(answers):
    """Mỗi câu AUTO: đáp án đúng = response duy nhất được chấm is_correct=True."""
    correct = {}
    for a in answers:
        if a['grading_method'] == 'AUTO' and a['is_correct'] is True:
            correct.setdefault(a['question_id'], set()).add(a['response_json'])
    out = {}
    for qid, vals in correct.items():
        assert len(vals) == 1, f'{qid} co {len(vals)} dap an dung khac nhau!'
        resp = json.loads(list(vals)[0])
        if 'selected' in resp:
            out[qid] = resp['selected']          # 'A' hoac ['A','B']
        elif 'text' in resp:
            out[qid] = resp['text']              # dien khuyet
    return out


def main():
    verify_only = '--verify' in sys.argv
    wb = load_workbook(XLSX, data_only=True)
    students_f = sheet_rows(wb, 'students')
    topics_f = sheet_rows(wb, 'topics')
    activities_f = sheet_rows(wb, 'learning_activities')
    questions_f = sheet_rows(wb, 'questions')
    act_q_f = sheet_rows(wb, 'activity_questions')
    assignments_f = sheet_rows(wb, 'assignments')
    asg_q_f = sheet_rows(wb, 'assignment_questions')
    targets_f = sheet_rows(wb, 'assignment_targets')
    attempts_f = sheet_rows(wb, 'attempts')
    answers_f = sheet_rows(wb, 'attempt_answers')
    reviews_f = sheet_rows(wb, 'answer_reviews')
    sessions_f = sheet_rows(wb, 'student_sessions')
    dim_f = sheet_rows(wb, 'dim_date')
    derived = derive_answers(answers_f)
    print('Dap an dung suy ra duoc:', derived)

    conn = connect()
    gv = conn.run("SELECT id FROM public.teachers WHERE username='gv001'")
    assert gv, 'chua co gv001'
    gv_id = str(gv[0][0])

    if not verify_only:
        for t in ('students', 'topics', 'attempts', 'student_sessions'):
            n = conn.run(f"SELECT count(*) FROM public.{t}")[0][0]
            assert n == 0, f'bang {t} chua trong ({n}) - dung lai'
        print('Kiem tra bang trong: OK')

        # ---- Lớp 7A4 ----
        class_id = str(uuid.uuid4())
        cls_name = str(students_f[0]['class_id']).replace('class_', '').upper()
        conn.run("""INSERT INTO public.classes (id, name, grade, school_year, teacher_id, status)
                    VALUES (:id, :name, 'Lớp 7', '2026-2027', :tid, 'active')""",
                 id=class_id, name=cls_name, tid=gv_id)

        # ---- Học sinh + student_states ----
        stu_map = {}
        sessions_by_student = {}
        for s in sessions_f:
            sessions_by_student.setdefault(s['student_id'], []).append(s)
        for s in students_f:
            sid = str(uuid.uuid4())
            stu_map[s['student_id']] = sid
            conn.run("""INSERT INTO public.students
                        (id, class_id, name, username, password_hash, student_code, status,
                         has_logged_in, last_active_at, created_at)
                        VALUES (:id, :cid, :name, :uname, :pw, :code, 'active', true, :last, :created)""",
                     id=sid, cid=class_id, name=s['full_name'],
                     uname=str(s['student_code']).lower(), pw=DEFAULT_STUDENT_PASSWORD,
                     code=s['student_code'], last=ts(s['last_login_at']), created=ts(s['created_at']))
            my_sessions = sessions_by_student.get(s['student_id'], [])
            login_days = sorted({dt_only(x['login_date']) for x in my_sessions})
            last_day = login_days[-1] if login_days else None
            streak = 0
            week_days = 0
            if last_day:
                cur = date.fromisoformat(last_day)
                dayset = set(login_days)
                while cur.isoformat() in dayset:
                    streak += 1
                    cur -= timedelta(days=1)
                monday = cur + timedelta(days=1)  # ngày đầu chuỗi liên tục gần nhất
                last_d = date.fromisoformat(last_day)
                week_start = last_d - timedelta(days=last_d.weekday())
                week_days = len([d for d in login_days
                                 if week_start <= date.fromisoformat(d) <= week_start + timedelta(days=6)])
            conn.run("""INSERT INTO public.student_states
                        (student_id, xp_today, xp_week, total_xp, active_seconds_today,
                         consecutive_days, last_attendance_date, completed_steps, topic_mastery,
                         unlocked_badge_ids, xp_logs, preferences, attendance_days_this_week,
                         attendance_history, question_results, flagged_topic_ids, current_progress)
                        VALUES (:sid, 0, 0, 0, 0, :streak, :lastd, '[]'::jsonb, '{}'::jsonb,
                                '[]'::jsonb, '[]'::jsonb, '{}'::jsonb, :week,
                                CAST(:hist AS jsonb), '[]'::jsonb, '[]'::jsonb, NULL)""",
                     sid=sid, streak=streak, lastd=last_day,
                     week=week_days, hist=json.dumps(login_days))
        print(f'Nap students: {len(stu_map)} + student_states')

        # ---- Chủ đề ----
        for i, t in enumerate(topics_f):
            conn.run("""INSERT INTO public.topics
                        (id, teacher_id, title, category, subject, grade, order_index,
                         video_payload, theory_payload, status)
                        VALUES (:id, :tid, :title, :cat, :subj, :grade, :ord,
                                CAST(:vp AS jsonb), CAST(:tp AS jsonb), 'published')""",
                     id=t['topic_id'], tid=gv_id, title=t['title'], cat=t['subject'],
                     subj=t['subject'], grade=t['grade'], ord=i + 1,
                     vp=json.dumps({'shortDesc': '', 'videoIds': [], 'class_ids': ['all'], 'version': 1}),
                     tp=json.dumps({'theoryId': ''}))
        print(f'Nap topics: {len(topics_f)}')

        # ---- Hoạt động học (luồng) ----
        for a in activities_f:
            conn.run("""INSERT INTO public.learning_activities
                        (id, topic_id, teacher_id, activity_type, title, flow_type, sequence_no, required_for_flow)
                        VALUES (:id, :topic, :tid, :atype, :title, :flow, :seq, :req)""",
                     id=a['activity_id'], topic=a['topic_id'], tid=gv_id,
                     atype=a['activity_type'], title=a['title'], flow=a['flow_type'],
                     seq=a['sequence_no'], req=boolean(a['required_for_flow']))
        print(f'Nap learning_activities: {len(activities_f)}')

        # ---- Câu hỏi (tích sẵn đáp án đúng cho câu AUTO) ----
        for q in questions_f:
            dbtype = TYPE_MAP[q['type']]
            ans = derived.get(q['question_id'])
            payload = {
                'id': q['question_id'], 'topicId': q['topic_id'], 'type': SHORT_TYPE[dbtype],
                'level': q['cognitive_level'], 'difficulty': DIFF_APP.get(q['difficulty'], 'TB'),
                'prompt': q['prompt'], 'options': [],
                'answer': ans if ans is not None else '',
                'explanation': '', 'points': num(q['max_score']),
                'isAiGenerated': False, 'isTeacherReviewed': True,
                'status': 'published', 'version': 1, 'class_ids': ['all'],
                'teacher_id': gv_id, 'max_score': num(q['max_score']),
                'grading_method': q['grading_method'],
                'created_at': datetime.now(TZ7).isoformat(), 'updated_at': datetime.now(TZ7).isoformat(),
            }
            conn.run("""INSERT INTO public.questions
                        (id, topic_id, teacher_id, type, level, difficulty, max_score, grading_method,
                         prompt, options, correct_answer, explanation, payload)
                        VALUES (:id, :topic, :tid, :type, :level, :diff, :maxs, :gm, :prompt,
                                '[]'::jsonb, CAST(:ans AS jsonb), '', CAST(:payload AS jsonb))""",
                     id=q['question_id'], topic=q['topic_id'], tid=gv_id, type=dbtype,
                     level=q['cognitive_level'], diff=q['difficulty'], maxs=num(q['max_score']),
                     gm=q['grading_method'], prompt=q['prompt'],
                     ans=json.dumps(ans, ensure_ascii=False) if ans is not None else None,
                     payload=json.dumps(payload, ensure_ascii=False))
        print(f'Nap questions: {len(questions_f)} (co dap an: {len(derived)})')

        # ---- Nối hoạt động - câu hỏi ----
        for r in act_q_f:
            conn.run("""INSERT INTO public.activity_questions (activity_id, question_id, position, points_override)
                        VALUES (:a, :q, :pos, :po)""",
                     a=r['activity_id'], q=r['question_id'], pos=r['position'], po=num(r['points_override']))
        print(f'Nap activity_questions: {len(act_q_f)}')

        # ---- Bài độc lập + câu hỏi + giao bài ----
        for a in assignments_f:
            conn.run("""INSERT INTO public.assignments
                        (id, title, activity_type, display_location, target_class_id, teacher_id,
                         published_at, due_at, is_timed, time_limit_minutes)
                        VALUES (:id, :title, :atype, :loc, :cid, :tid, :pub, :due, :timed, :limit)""",
                     id=a['assignment_id'], title=a['title'], atype=a['activity_type'],
                     loc=a['display_location'], cid=class_id, tid=gv_id,
                     pub=ts(a['published_at']), due=ts(a['due_at']),
                     timed=boolean(a['is_timed']), limit=a['time_limit_minutes'])
        for r in asg_q_f:
            conn.run("""INSERT INTO public.assignment_questions (assignment_id, question_id, position)
                        VALUES (:a, :q, :pos)""",
                     a=r['assignment_id'], q=r['question_id'], pos=r['position'])
        for t in targets_f:
            conn.run("""INSERT INTO public.assignment_targets
                        (target_id, assignment_id, student_id, assigned_at, due_at, submission_status, first_submitted_at)
                        VALUES (:id, :a, :sid, :assigned, :due, :status, :first)""",
                     id=t['target_id'], a=t['assignment_id'], sid=stu_map[t['student_id']],
                     assigned=ts(t['assigned_at']), due=ts(t['due_at']),
                     status=t['submission_status'], first=ts(t['first_submitted_at']))
        print(f"Nap assignments: {len(assignments_f)} + {len(asg_q_f)} cau + {len(targets_f)} targets")

        # ---- Lượt làm bài + câu trả lời + phiếu chấm ----
        for a in attempts_f:
            conn.run("""INSERT INTO public.attempts
                        (attempt_id, student_id, assignment_id, learning_activity_id, attempt_no,
                         status, submitted_at, submitted_date, score_awarded, max_score, grading_status, grading_mode)
                        VALUES (:id, :sid, :asg, :act, :no, :status, :sub, :subd, :score, :maxs, :gs, :gm)""",
                     id=a['attempt_id'], sid=stu_map[a['student_id']],
                     asg=a['assignment_id'], act=a['learning_activity_id'], no=a['attempt_no'],
                     status=a['status'], sub=ts(a['submitted_at']), subd=dt_only(a['submitted_date']),
                     score=num(a['score_awarded']), maxs=num(a['max_score']),
                     gs=a['grading_status'], gm=a['grading_mode'])
        for a in answers_f:
            conn.run("""INSERT INTO public.attempt_answers
                        (answer_id, attempt_id, question_id, response_json, is_correct, score_awarded,
                         max_score, grading_status, grading_method, response_time_seconds)
                        VALUES (:id, :att, :q, CAST(:resp AS jsonb), :corr, :score, :maxs, :gs, :gm, :rtime)""",
                     id=a['answer_id'], att=a['attempt_id'], q=a['question_id'],
                     resp=a['response_json'] if isinstance(a['response_json'], str) else json.dumps(a['response_json']),
                     corr=a['is_correct'] if isinstance(a['is_correct'], bool) else None,
                     score=num(a['score_awarded']), maxs=num(a['max_score']),
                     gs=a['grading_status'], gm=a['grading_method'], rtime=a['response_time_seconds'])
        for r in reviews_f:
            rubric = None
            if r.get('rubric_text'):
                try:
                    rubric = json.loads(r['rubric_text'])
                except (TypeError, json.JSONDecodeError):
                    rubric = {'text': str(r['rubric_text'])}
            conn.run("""INSERT INTO public.answer_reviews
                        (review_id, answer_id, reviewer_id, score_awarded, review_status, feedback,
                         rubric_scores, reviewed_at, created_at)
                        VALUES (:id, :ans, :rev, :score, :status, :fb, CAST(:rubric AS jsonb), :at, :created)""",
                     id=r['review_id'], ans=r['answer_id'], rev=gv_id,
                     score=num(r['score_awarded']), status=r['review_status'], fb=r['feedback'],
                     rubric=json.dumps(rubric, ensure_ascii=False) if rubric is not None else None,
                     at=ts(r['reviewed_at']), created=ts(r['created_at']))
        print(f"Nap attempts: {len(attempts_f)} + answers: {len(answers_f)} + reviews: {len(reviews_f)}")

        # ---- Phiên đăng nhập + bảng ngày ----
        for s in sessions_f:
            conn.run("""INSERT INTO public.student_sessions
                        (session_id, student_id, login_at, logout_at, login_date,
                         session_duration_minutes, device_type, session_status)
                        VALUES (:id, :sid, :login, :logout, :ldate, :dur, :dev, :status)""",
                     id=s['session_id'], sid=stu_map[s['student_id']],
                     login=ts(s['login_at']), logout=ts(s['logout_at']), ldate=dt_only(s['login_date']),
                     dur=s['session_duration_minutes'], dev=s['device_type'], status=s['session_status'])
        for drow in dim_f:
            conn.run("""INSERT INTO public.dim_date
                        (date, year, month_number, month_name, year_month, week_start, day_of_week, is_weekend)
                        VALUES (:d, :y, :m, :mname, :ym, :ws, :dow, :we)""",
                     d=dt_only(drow['date']), y=drow['year'], m=drow['month_number'],
                     mname=drow['month_name'], ym=drow['year_month'], ws=dt_only(drow['week_start']),
                     dow=drow['day_of_week'], we=boolean(drow['is_weekend']))
        print(f"Nap sessions: {len(sessions_f)} + dim_date: {len(dim_f)}")
        conn.commit()
        print('DA COMMIT - NAP XONG')

    # ================= KIỂM CHỨNG (bộ số vàng) =================
    print('\n===== KIEM CHUNG =====')
    tables = {'students': 10, 'topics': 2, 'learning_activities': 16, 'questions': 10,
              'activity_questions': 30, 'assignments': 3, 'assignment_questions': 11,
              'assignment_targets': 30, 'attempts': 54, 'attempt_answers': 223,
              'answer_reviews': 44, 'student_sessions': 32, 'dim_date': 61}
    ok = True
    for t, want in tables.items():
        got = conn.run(f"SELECT count(*) FROM public.{t}")[0][0]
        flag = 'OK' if got == want else 'LECH!'
        if got != want: ok = False
        print(f'{t}: {got}/{want} {flag}')
    checks = [
        ("attempts theo bai doc lap", "SELECT count(*) FROM attempts WHERE assignment_id IS NOT NULL", 28),
        ("attempts theo luong hoc", "SELECT count(*) FROM attempts WHERE learning_activity_id IS NOT NULL", 26),
        ("attempts vi pham XOR", "SELECT count(*) FROM attempts WHERE (assignment_id IS NOT NULL) = (learning_activity_id IS NOT NULL)", 0),
        ("cau cho cham (answers PENDING)", "SELECT count(*) FROM attempt_answers WHERE grading_status='PENDING'", 39),
        ("phieu cho cham (reviews PENDING)", "SELECT count(*) FROM answer_reviews WHERE review_status='PENDING'", 39),
        ("targets SUBMITTED", "SELECT count(*) FROM assignment_targets WHERE submission_status='SUBMITTED'", 24),
        ("targets NOT_SUBMITTED", "SELECT count(*) FROM assignment_targets WHERE submission_status='NOT_SUBMITTED'", 6),
        ("attempts chua co diem (PENDING_TEACHER_REVIEW)", "SELECT count(*) FROM attempts WHERE grading_status='PENDING_TEACHER_REVIEW'", 3),
        ("HS dang nhap 7 ngay cuoi cua du lieu (15-21/10)", "SELECT count(DISTINCT student_id) FROM student_sessions WHERE login_date BETWEEN '2026-10-15' AND '2026-10-21'", 7),
        ("HS dang nhap 7 ngay lien truoc (08-14/10)", "SELECT count(DISTINCT student_id) FROM student_sessions WHERE login_date BETWEEN '2026-10-08' AND '2026-10-14'", 6),
        ("cau hoi co dap an dung", "SELECT count(*) FROM questions WHERE correct_answer IS NOT NULL", 6),
    ]
    for label, sql, want in checks:
        got = conn.run(sql)[0][0]
        flag = 'OK' if got == want else 'LECH!'
        if got != want: ok = False
        print(f'{label}: {got} (ky vong {want}) {flag}')
    # Đối soát điểm: attempts đã chấm phải bằng tổng điểm các câu
    rows = conn.run("""SELECT a.attempt_id, a.score_awarded,
        (SELECT COALESCE(SUM(ans.score_awarded),0) FROM attempt_answers ans WHERE ans.attempt_id = a.attempt_id)
        FROM attempts a WHERE a.grading_status = 'GRADED'""")
    bad = [r for r in rows if r[1] is None or abs(float(r[1]) - float(r[2])) > 1e-9]
    print(f'Doi soat diem attempts GRADED: {len(rows)} attempts, lech: {len(bad)}', 'OK' if not bad else 'LECH!')
    if bad: ok = False
    # last_active_at khớp phiên đăng nhập gần nhất
    rows = conn.run("""SELECT s.username FROM students s
        WHERE s.last_active_at IS DISTINCT FROM (SELECT MAX(ss.login_at) FROM student_sessions ss WHERE ss.student_id = s.id)""")
    print(f'HS lech last_active_at vs sessions: {len(rows)}', 'OK' if not rows else 'LECH!')
    if rows: ok = False
    print('\nTONG KET KIEM CHUNG:', 'DAT TOAN BO' if ok else 'CON DIEM LECH - XEM LAI')
    conn.close()


if __name__ == '__main__':
    main()
